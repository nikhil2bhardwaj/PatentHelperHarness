import os
import sys
import json
import uuid
import asyncio
import re
import tempfile
from typing import Dict, Optional, List
from pydantic import BaseModel
from fastapi import FastAPI, HTTPException, BackgroundTasks, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, PlainTextResponse, Response

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from harness.config import PHASE_NAMES, USE_JSON_OUTPUT, GROQ_MODEL, OPENAI_MODEL, OUTPUTS_DIR
from harness.api_client import APIClient
from harness.mock_client import MockAPIClient
from harness.output_store import OutputStore
from harness.prompt_loader import load_system_prompt, load_all_prompts, substitute_variables
from harness.utils import count_tokens, compress_phase_output, build_historical_context, build_cumulative_context

app = FastAPI(title="Patent Helper Harness API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.middleware("http")
async def add_no_cache_header(request, call_next):
    response = await call_next(request)
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response

sessions: Dict[str, dict] = {}


class StartSessionRequest(BaseModel):
    user_input: str = ""
    provider: str = "groq"
    model: Optional[str] = None
    mock: bool = False


class GateRequest(BaseModel):
    action: str  # "accept", "edit", "feedback", "exit", "rerun"
    edited_content: Optional[str] = None
    feedback: Optional[str] = None


class ResumeSessionRequest(BaseModel):
    session_dir: str


class RunPhaseRequest(BaseModel):
    referenced_phases: List[int] = []


@app.get("/api/health")
async def health():
    return {"status": "ok"}


@app.get("/api/phases")
async def get_phases():
    return {
        "phases": [
            {"phase": p, "name": PHASE_NAMES.get(p, f"Phase {p}")}
            for p in range(13)
        ]
    }


@app.post("/api/sessions")
async def start_session(req: StartSessionRequest):
    session_id = str(uuid.uuid4())[:8]
    output_store = OutputStore()

    if req.mock:
        api_client = MockAPIClient(model=req.model)
    else:
        api_client = APIClient(model=req.model, provider=req.provider)

    prompts = load_all_prompts()
    system_prompt = load_system_prompt()

    sessions[session_id] = {
        "api_client": api_client,
        "output_store": output_store,
        "prompts": prompts,
        "system_prompt": system_prompt,
        "user_input": req.user_input,
        "provider": req.provider,
        "phase_outputs": {},
        "feedback": {},
        "current_phase": 0,
        "status": "idle",  # idle, running, awaiting_gate, completed, exited
        "last_output": None,
    }

    return {
        "session_id": session_id,
        "session_dir": output_store.session_dir,
        "phases": [
            {"phase": p, "name": PHASE_NAMES.get(p, f"Phase {p}")}
            for p in range(13)
        ],
    }


@app.post("/api/sessions/{session_id}/run/{phase}")
async def run_phase(session_id: str, phase: int, req: RunPhaseRequest = None):
    if session_id not in sessions:
        raise HTTPException(status_code=404, detail="Session not found")
    if phase < 0 or phase > 12:
        raise HTTPException(status_code=400, detail=f"Invalid phase {phase}. Must be between 0 and 12.")

    sess = sessions[session_id]
    sess["status"] = "running"
    sess["current_phase"] = phase

    try:
        prompt = sess["prompts"][phase]
        values = _build_variable_values(sess, phase)
        user_prompt = substitute_variables(prompt.template, values)

        if phase in sess["feedback"]:
            user_prompt = f"[ADDITIONAL USER FEEDBACK]\n{sess['feedback'][phase]}\n[/ADDITIONAL USER FEEDBACK]\n\n{user_prompt}"

        referenced_phases = req.referenced_phases if req else []
        if referenced_phases:
            ref_parts = []
            for rp in referenced_phases:
                rp_output = sess["phase_outputs"].get(rp)
                if rp_output is None:
                    rp_output = sess["output_store"].load_phase_output(rp)
                if rp_output:
                    ref_name = PHASE_NAMES.get(rp, f"Phase {rp}")
                    compressed = compress_phase_output(rp_output, next_phase=phase, model=sess["api_client"].model)
                    ref_parts.append(f"[REFERENCED OUTPUT - Phase {rp}: {ref_name}]\n{compressed}\n[/REFERENCED OUTPUT]")
            if ref_parts:
                ref_block = "\n\n".join(ref_parts)
                user_prompt = f"{ref_block}\n\n{user_prompt}"

        token_count = count_tokens(sess["system_prompt"]) + count_tokens(user_prompt)

        loop = asyncio.get_event_loop()
        output = await loop.run_in_executor(
            None,
            sess["api_client"].chat,
            sess["system_prompt"],
            user_prompt,
            USE_JSON_OUTPUT,
        )

        sess["output_store"].save_phase_output(phase, output)
        sess["last_output"] = output
        sess["status"] = "awaiting_gate"

        parsed = None
        try:
            parsed = json.loads(output)
        except (json.JSONDecodeError, TypeError):
            pass

        return {
            "phase": phase,
            "phase_name": prompt.name,
            "output": output,
            "parsed": parsed,
            "estimated_input_tokens": token_count,
            "status": "awaiting_gate",
        }
    except Exception as e:
        sess["status"] = "error"
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/sessions/{session_id}/gate")
async def gate_decision(session_id: str, req: GateRequest):
    if session_id not in sessions:
        raise HTTPException(status_code=404, detail="Session not found")

    sess = sessions[session_id]
    phase = sess["current_phase"]
    output = sess["last_output"]

    if req.action == "accept":
        sess["phase_outputs"][phase] = output
        sess["status"] = "idle"
        _save_state(sess)
        next_phase = phase + 1 if phase < 12 else None
        return {"action": "accepted", "next_phase": next_phase, "completed": next_phase is None}

    elif req.action == "edit":
        if not req.edited_content:
            raise HTTPException(status_code=400, detail="edited_content required for edit action")
        sess["output_store"].save_phase_output(phase, req.edited_content, edited=True)
        sess["phase_outputs"][phase] = req.edited_content
        sess["status"] = "idle"
        _save_state(sess)
        next_phase = phase + 1 if phase < 12 else None
        return {"action": "edited", "next_phase": next_phase, "completed": next_phase is None}

    elif req.action == "feedback":
        if not req.feedback:
            raise HTTPException(status_code=400, detail="feedback required for feedback action")
        sess["feedback"][phase + 1] = req.feedback
        sess["output_store"].save_feedback(phase, req.feedback)
        sess["phase_outputs"][phase] = output
        sess["status"] = "idle"
        _save_state(sess)
        next_phase = phase + 1 if phase < 12 else None
        return {"action": "feedback_saved", "next_phase": next_phase, "completed": next_phase is None}

    elif req.action == "exit":
        sess["phase_outputs"][phase] = output
        sess["status"] = "exited"
        _save_state(sess)
        return {"action": "saved", "session_dir": sess["output_store"].session_dir}

    elif req.action == "rerun":
        sess["status"] = "idle"
        return {"action": "rerun", "next_phase": phase}

    else:
        raise HTTPException(status_code=400, detail=f"Unknown action: {req.action}")


@app.get("/api/sessions/{session_id}")
async def get_session(session_id: str):
    if session_id not in sessions:
        raise HTTPException(status_code=404, detail="Session not found")

    sess = sessions[session_id]
    return {
        "session_id": session_id,
        "session_dir": sess["output_store"].session_dir,
        "current_phase": sess["current_phase"],
        "status": sess["status"],
        "completed_phases": sorted(sess["phase_outputs"].keys()),
        "user_input": sess["user_input"],
        "model": sess["api_client"].model,
    }


@app.get("/api/sessions/{session_id}/cumulative")
async def get_cumulative(session_id: str):
    if session_id not in sessions:
        raise HTTPException(status_code=404, detail="Session not found")

    sess = sessions[session_id]
    content = sess["output_store"].load_cumulative()
    if not content:
        return {"content": ""}
    return {"content": content}


@app.get("/api/sessions/{session_id}/phase/{phase}")
async def get_phase_output(session_id: str, phase: int):
    if session_id not in sessions:
        raise HTTPException(status_code=404, detail="Session not found")

    sess = sessions[session_id]
    output = sess["output_store"].load_phase_output(phase)
    if not output:
        raise HTTPException(status_code=404, detail=f"Phase {phase} output not found")

    parsed = None
    try:
        parsed = json.loads(output)
    except (json.JSONDecodeError, TypeError):
        pass

    return {"phase": phase, "output": output, "parsed": parsed}


def _build_variable_values(sess: dict, phase: int) -> Dict[str, str]:
    values: Dict[str, str] = {}

    if phase == 0:
        values["USER_INPUT"] = sess["user_input"] if sess["user_input"] else "(No user input provided. Proceed in Autonomous mode.)"

    if phase >= 1:
        prev_output = sess["phase_outputs"].get(phase - 1)
        if prev_output is None:
            prev_output = sess["output_store"].load_phase_output(phase - 1)
            if prev_output:
                sess["phase_outputs"][phase - 1] = prev_output
        if prev_output:
            from harness.config import COMPRESS_PHASE_OUTPUTS
            if COMPRESS_PHASE_OUTPUTS:
                values[f"PHASE_{phase - 1}_OUTPUT"] = compress_phase_output(
                    prev_output, next_phase=phase, model=sess["api_client"].model
                )
            else:
                values[f"PHASE_{phase - 1}_OUTPUT"] = prev_output

    if phase == 11:
        historical = build_historical_context(
            sess["phase_outputs"], phases=[5, 6, 7, 8, 9],
            model=sess["api_client"].model,
        )
        values["HISTORICAL_CONTEXT"] = historical

    if phase == 12:
        cumulative = build_cumulative_context(
            sess["phase_outputs"], phases=[4, 5, 6, 7, 8, 9, 10],
            model=sess["api_client"].model,
        )
        values["CUMULATIVE_CONTEXT"] = cumulative

    return values


def _save_state(sess: dict):
    state = {
        "current_phase": sess["current_phase"],
        "completed": False,
        "completed_phases": sorted(sess["phase_outputs"].keys()),
        "user_input": sess["user_input"],
        "model": sess["api_client"].model,
        "provider": sess.get("provider", "groq"),
    }
    sess["output_store"].save_session_state(state)


@app.get("/api/sessions/list")
async def list_saved_sessions():
    from harness.config import OUTPUTS_DIR
    session_dirs = OutputStore.list_sessions()
    result = []
    for sdir in session_dirs:
        store = OutputStore(session_dir=sdir)
        state = store.load_session_state()
        if state:
            result.append({
                "session_dir": sdir,
                "session_name": os.path.basename(sdir),
                "current_phase": state.get("current_phase", 0),
                "completed_phases": state.get("completed_phases", []),
                "user_input": state.get("user_input", ""),
                "model": state.get("model", ""),
                "provider": state.get("provider", "groq"),
            })
        else:
            # No state file - check if any phase outputs exist
            completed = []
            for p in range(13):
                if store.load_phase_output(p):
                    completed.append(p)
            if completed:
                result.append({
                    "session_dir": sdir,
                    "session_name": os.path.basename(sdir),
                    "current_phase": max(completed) + 1 if max(completed) < 12 else 12,
                    "completed_phases": completed,
                    "user_input": "",
                    "model": "",
                    "provider": "groq",
                })
    return {"sessions": result}


@app.post("/api/sessions/resume")
async def resume_session(req: ResumeSessionRequest):
    session_dir = req.session_dir
    abs_session = os.path.abspath(session_dir)
    abs_outputs = os.path.abspath(OUTPUTS_DIR)

    # Security check: prevent directory traversal outside OUTPUTS_DIR
    if not abs_session.startswith(abs_outputs):
        raise HTTPException(status_code=400, detail="Invalid session path: must reside within outputs directory")

    if not os.path.isdir(abs_session):
        raise HTTPException(status_code=404, detail="Session directory not found")

    output_store = OutputStore(session_dir=session_dir)
    state = output_store.load_session_state()

    provider = "groq"
    model = None
    user_input = ""
    completed_phases = []
    current_phase = 0

    if state:
        provider = state.get("provider", "groq")
        model = state.get("model")
        user_input = state.get("user_input", "")
        completed_phases = state.get("completed_phases", [])
        current_phase = state.get("current_phase", 0)

    # Load all completed phase outputs back into memory
    phase_outputs = {}
    for p in completed_phases:
        output = output_store.load_phase_output(p)
        if output:
            phase_outputs[p] = output

    # Determine next phase to run
    if completed_phases:
        next_phase = max(completed_phases) + 1
        if next_phase > 12:
            next_phase = 12
        current_phase = next_phase

    # Create API client with saved provider/model
    api_client = APIClient(model=model, provider=provider)

    prompts = load_all_prompts()
    system_prompt = load_system_prompt()

    session_id = str(uuid.uuid4())[:8]
    sessions[session_id] = {
        "api_client": api_client,
        "output_store": output_store,
        "prompts": prompts,
        "system_prompt": system_prompt,
        "user_input": user_input,
        "provider": provider,
        "phase_outputs": phase_outputs,
        "feedback": {},
        "current_phase": current_phase,
        "status": "idle",
        "last_output": None,
    }

    return {
        "session_id": session_id,
        "session_dir": session_dir,
        "current_phase": current_phase,
        "completed_phases": sorted(phase_outputs.keys()),
        "user_input": user_input,
        "provider": provider,
        "model": api_client.model,
        "phases": [
            {"phase": p, "name": PHASE_NAMES.get(p, f"Phase {p}")}
            for p in range(13)
        ],
    }


def _json_to_markdown(data, depth=0, phase_name=""):
    lines = []
    if depth == 0 and phase_name:
        lines.append(f"# Phase Output: {phase_name}\n")
    if isinstance(data, dict):
        for key, val in data.items():
            formatted_key = re.sub(r'^\d+_', '', str(key)).replace('_', ' ').title()
            if isinstance(val, dict):
                lines.append(f"{'#' * (depth + 2)} {formatted_key}\n")
                lines.append(_json_to_markdown(val, depth + 1))
            elif isinstance(val, list):
                lines.append(f"{'#' * (depth + 2)} {formatted_key}\n")
                for item in val:
                    if isinstance(item, dict):
                        lines.append(_json_to_markdown(item, depth + 2))
                    else:
                        lines.append(f"- {item}")
                lines.append("")
            else:
                lines.append(f"**{formatted_key}:** {val}\n")
    elif isinstance(data, list):
        for item in data:
            if isinstance(item, dict):
                lines.append(_json_to_markdown(item, depth + 1))
            else:
                lines.append(f"- {item}")
    else:
        lines.append(str(data))
    return "\n".join(lines)


def _generate_pdf(markdown_text: str, phase_name: str) -> bytes:
    from fpdf import FPDF
    pdf = FPDF(format='A4')
    pdf.set_margins(15, 15, 15)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    epw = pdf.epw

    # Document Header
    pdf.set_font("helvetica", 'B', 16)
    pdf.set_x(pdf.l_margin)
    safe_phase = phase_name.encode('latin-1', 'replace').decode('latin-1')
    pdf.multi_cell(epw, 8, f"Phase Output: {safe_phase}")
    pdf.ln(3)

    for line in markdown_text.split('\n'):
        line_str = line.strip()
        if not line_str:
            pdf.ln(2)
            continue

        pdf.set_x(pdf.l_margin)
        if line_str.startswith('# '):
            pdf.set_font("helvetica", 'B', 14)
            safe_text = line_str[2:].encode('latin-1', 'replace').decode('latin-1')
            pdf.multi_cell(epw, 7, safe_text)
            pdf.ln(1)
        elif line_str.startswith('## '):
            pdf.set_font("helvetica", 'B', 12)
            safe_text = line_str[3:].encode('latin-1', 'replace').decode('latin-1')
            pdf.multi_cell(epw, 6, safe_text)
            pdf.ln(1)
        elif line_str.startswith('### '):
            pdf.set_font("helvetica", 'B', 11)
            safe_text = line_str[4:].encode('latin-1', 'replace').decode('latin-1')
            pdf.multi_cell(epw, 6, safe_text)
            pdf.ln(1)
        elif line_str.startswith('- '):
            pdf.set_font("helvetica", '', 10)
            safe_text = ("- " + line_str[2:]).encode('latin-1', 'replace').decode('latin-1')
            pdf.multi_cell(epw, 5, safe_text)
        else:
            clean = re.sub(r'\*\*(.+?)\*\*', r'\1', line_str)
            safe_text = clean.encode('latin-1', 'replace').decode('latin-1')
            pdf.set_font("helvetica", '', 10)
            pdf.multi_cell(epw, 5, safe_text)

    return bytes(pdf.output())


def _generate_docx(markdown_text: str, phase_name: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, Inches
    import io

    doc = Document()
    doc.add_heading(f"Phase Output: {phase_name}", 0)

    for line in markdown_text.split('\n'):
        if line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
        elif line.strip() == '':
            pass
        else:
            clean = re.sub(r'\*\*(.+?)\*\*', r'\1', line)
            doc.add_paragraph(clean)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@app.get("/api/sessions/{session_id}/phase/{phase}/export")
async def export_phase_output(session_id: str, phase: int, format: str = Query("md", pattern="^(md|pdf|docx)$")):
    if session_id not in sessions:
        raise HTTPException(status_code=404, detail="Session not found")

    sess = sessions[session_id]
    output = sess["output_store"].load_phase_output(phase)
    if not output:
        raise HTTPException(status_code=404, detail=f"Phase {phase} output not found")

    phase_name = PHASE_NAMES.get(phase, f"Phase {phase}")

    parsed = None
    try:
        parsed = json.loads(output)
    except (json.JSONDecodeError, TypeError):
        pass

    if parsed:
        markdown = _json_to_markdown(parsed, phase_name=phase_name)
    else:
        markdown = f"# Phase Output: {phase_name}\n\n{output}"

    safe_name = re.sub(r'[^a-zA-Z0-9_\-]', '_', phase_name)
    filename = f"phase_{phase}_{safe_name}"

    if format == "md":
        return PlainTextResponse(
            content=markdown,
            media_type="text/markdown",
            headers={"Content-Disposition": f"attachment; filename={filename}.md"}
        )
    elif format == "pdf":
        pdf_bytes = _generate_pdf(markdown, phase_name)
        return Response(
            content=pdf_bytes,
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename={filename}.pdf"}
        )
    elif format == "docx":
        docx_bytes = _generate_docx(markdown, phase_name)
        return Response(
            content=docx_bytes,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename={filename}.docx"}
        )


frontend_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "frontend", "dist")
if os.path.isdir(frontend_dir):
    app.mount("/", StaticFiles(directory=frontend_dir, html=True), name="frontend")
